"""Pruebas de la lógica que no depende de Streamlit."""

from datetime import date

import pytest
from docx import Document

from curaduria import fechas, numeros, plantilla
from curaduria.genero_nombres import sugerir_genero
from curaduria.gramatica import Concordancia, Titular
from curaduria.resolucion import Inmueble, Resolucion


# --- Números ---------------------------------------------------------------

@pytest.mark.parametrize("valor,esperado", [
    (0, "cero"), (1, "uno"), (15, "quince"), (21, "veintiuno"),
    (31, "treinta y uno"), (100, "cien"), (101, "ciento uno"),
    (120, "ciento veinte"), (999, "novecientos noventa y nueve"),
    (1000, "mil"), (1001, "mil uno"), (2026, "dos mil veintiséis"),
    (15000, "quince mil"), (1_000_000, "un millón"),
    (2_500_000, "dos millones quinientos mil"),
])
def test_entero_a_letras(valor, esperado):
    assert numeros.entero_a_letras(valor) == esperado


def test_entero_fuera_de_rango():
    with pytest.raises(ValueError):
        numeros.entero_a_letras(1_000_000_000)


def test_decimal_a_letras():
    assert numeros.decimal_a_letras(120.5) == "ciento veinte con cincuenta"
    assert numeros.decimal_a_letras(0.05) == "cero con cinco"


def test_decimal_redondeo_desborda():
    assert numeros.decimal_a_letras(1.999) == "dos"


def test_formato_cifra_colombiano():
    assert numeros.formato_cifra(1234567.5) == "1.234.567,50"


def test_en_letras_y_cifra():
    assert numeros.en_letras_y_cifra(120.5, "m2") == (
        "ciento veinte con cincuenta m2 (120,50 m2)")


def test_concordar_uno():
    assert numeros.concordar_uno("veintiuno", "F") == "veintiuna"
    assert numeros.concordar_uno("veintiuno", "M") == "veintiún"
    assert numeros.concordar_uno("treinta y uno", "M") == "treinta y un"


# --- Fechas ----------------------------------------------------------------

def test_fechas():
    f = date(2026, 10, 15)
    assert fechas.fecha_corta(f) == "15/10/2026"
    assert fechas.fecha_larga(f) == "15 de octubre de 2026"
    assert fechas.fecha_en_letras(f) == (
        "quince (15) de octubre de dos mil veintiséis (2026)")


def test_etiquetas_fecha_vacia():
    assert fechas.etiquetas_fecha("fecha", None)["fecha_letras"] == ""


# --- Gramática -------------------------------------------------------------

def test_un_titular_femenino():
    c = Concordancia([Titular("MARÍA GÓMEZ", "F", "CC", "41000000")])
    assert c.palabra("articulo") == "la"
    assert c.palabra("propietario") == "propietaria"
    assert c.palabra("identificado") == "identificada"
    assert c.palabra("verbo_solicitar") == "solicita"
    assert c.frase_calidad() == "en calidad de propietaria"


def test_dos_titulares_mixto_va_a_masculino_plural():
    c = Concordancia([
        Titular("JUAN PÉREZ", "M", "CC", "79000000"),
        Titular("MARÍA GÓMEZ", "F", "CC", "41000000"),
    ])
    assert c.mixto
    assert c.palabra("articulo") == "los"
    assert c.palabra("propietario") == "propietarios"
    assert c.palabra("verbo_podra") == "podrán"
    assert c.nombres() == "JUAN PÉREZ y MARÍA GÓMEZ"


def test_tres_titulares_femeninas():
    c = Concordancia([
        Titular("ANA DÍAZ", "F", "CC", "1"),
        Titular("LUZ RUIZ", "F", "CC", "2"),
        Titular("ROSA LIMA", "F", "CC", "3"),
    ])
    assert c.palabra("articulo") == "las"
    assert c.palabra("tratamiento") == "señoras"
    assert c.nombres() == "ANA DÍAZ, LUZ RUIZ y ROSA LIMA"


def test_persona_juridica_se_identifica_en_femenino():
    t = Titular("CONSTRUCTORA XYZ S.A.S.", "M", "NIT", "900123456-1")
    assert "identificada con Número de Identificación Tributaria" in \
        t.frase_identificacion()


def test_etiquetas_incluyen_version_capitalizada():
    c = Concordancia([Titular("MARÍA GÓMEZ", "F")])
    e = c.etiquetas()
    assert e["articulo"] == "la" and e["Articulo"] == "La"


def test_flexionar_libre():
    from curaduria.gramatica import flexionar_libre
    assert flexionar_libre("arquitecto", "F", "S") == "arquitecta"
    assert flexionar_libre("arquitecto", "M", "P") == "arquitectos"


# --- Sugerencia de género --------------------------------------------------

@pytest.mark.parametrize("nombre,esperado", [
    ("MARÍA FERNANDA GÓMEZ", "F"),
    ("JUAN CARLOS PÉREZ", "M"),
    ("JOSÉ MARÍA RESTREPO", "M"),   # gana el primer nombre reconocido
    ("LUZ MARINA TORRES", "F"),
    ("YULIANA ESCOBAR", "F"),       # por terminación
    ("", "M"),
])
def test_sugerir_genero(nombre, esperado):
    assert sugerir_genero(nombre) == esperado


# --- Plantilla -------------------------------------------------------------

def _doc_con(texto_runs):
    doc = Document()
    parrafo = doc.add_paragraph()
    for fragmento in texto_runs:
        parrafo.add_run(fragmento)
    return doc


def test_reemplazo_simple():
    doc = _doc_con(["Hola {{nombre}}, bienvenido."])
    plantilla.rellenar(doc, {"nombre": "ANA"})
    assert doc.paragraphs[0].text == "Hola ANA, bienvenido."


def test_reemplazo_de_etiqueta_partida_en_varios_runs():
    doc = _doc_con(["Hola {{nom", "bre}}", " y adiós"])
    plantilla.rellenar(doc, {"nombre": "ANA"})
    assert doc.paragraphs[0].text == "Hola ANA y adiós"


def test_varias_etiquetas_en_un_parrafo():
    doc = _doc_con(["{{a}}-{{b}}-{{a}}"])
    plantilla.rellenar(doc, {"a": "1", "b": "2"})
    assert doc.paragraphs[0].text == "1-2-1"


def test_etiqueta_con_espacios_internos():
    doc = _doc_con(["X {{ nombre }} Y"])
    plantilla.rellenar(doc, {"nombre": "ANA"})
    assert doc.paragraphs[0].text == "X ANA Y"


def test_etiqueta_sin_dato_se_reporta_y_se_conserva():
    doc = _doc_con(["Hola {{desconocida}}"])
    faltantes = plantilla.rellenar(doc, {"nombre": "ANA"})
    assert faltantes == {"desconocida"}
    assert doc.paragraphs[0].text == "Hola {{desconocida}}"


def test_reemplazo_conserva_formato_del_run():
    doc = Document()
    parrafo = doc.add_paragraph()
    run = parrafo.add_run("{{nombre}}")
    run.bold = True
    plantilla.rellenar(doc, {"nombre": "ANA"})
    assert doc.paragraphs[0].runs[0].bold is True
    assert doc.paragraphs[0].text == "ANA"


def test_reemplazo_en_tablas_encabezado_y_pie():
    doc = Document()
    tabla = doc.add_table(rows=1, cols=1)
    tabla.cell(0, 0).paragraphs[0].add_run("Mat: {{matricula}}")
    doc.sections[0].header.paragraphs[0].add_run("{{curaduria}}")
    doc.sections[0].footer.paragraphs[0].add_run("Res. {{resolucion_numero}}")
    plantilla.rellenar(doc, {"matricula": "050-123", "curaduria": "CU1",
                             "resolucion_numero": "0456"})
    assert tabla.cell(0, 0).text == "Mat: 050-123"
    assert doc.sections[0].header.paragraphs[0].text == "CU1"
    assert doc.sections[0].footer.paragraphs[0].text == "Res. 0456"


def test_listar_etiquetas_del_documento():
    doc = _doc_con(["{{b}} {{a}} {{b}}"])
    assert plantilla.etiquetas_del_documento(doc) == ["b", "a"]


# --- Resolución ------------------------------------------------------------

def _resolucion_valida():
    from curaduria.catalogo import Solicitud
    from curaduria.profesionales import Profesional
    from curaduria.resolucion import Despacho
    return Resolucion(
        numero="73001-2-26-0456",
        fecha=date(2026, 10, 15),
        radicado="2600789",
        fecha_radicado=date(2026, 3, 10),
        consecutivo_expediente="54",
        solicitud=Solicitud("construccion", ["obra_nueva"],
                            ["vivienda_unifamiliar"]),
        despacho=Despacho(curador="SANDRA DEL PILAR PARDO SUAREZ",
                          sustanciador="GLORIA INES LOZANO TRUJILLO",
                          sustanciador_iniciales="gilt /S"),
        titulares=[Titular("MARÍA GÓMEZ", "F", "CC", "41000000")],
        inmueble=Inmueble(direccion="CL 10 # 5-20",
                          matricula_inmobiliaria="350-123456",
                          ficha_catastral="01-09-1014-0013-000",
                          area_lote=120.5, area_construida=85.0, pisos=2),
        profesionales=[
            Profesional("proyectista", "FERNANDA GARCÍA", "F", "93383328",
                        "25700-55387"),
            Profesional("constructor", "NOLBERTO GONZÁLEZ", "M", "93377891",
                        "A16472015"),
        ],
    )


def test_resolucion_valida_sin_errores():
    assert _resolucion_valida().errores() == []


def test_errores_detectados():
    errores = Resolucion().errores()
    assert any("número de la resolución" in e for e in errores)
    assert any("al menos un titular" in e for e in errores)
    assert any("matrícula inmobiliaria" in e for e in errores)
    assert any("quien firma" in e for e in errores)


def test_error_documento_repetido():
    r = _resolucion_valida()
    r.titulares.append(Titular("OTRA PERSONA", "F", "CC", "41000000"))
    assert any("repetido" in e for e in r.errores())


def test_error_radicado_posterior():
    r = _resolucion_valida()
    r.fecha_radicado = date(2026, 12, 1)
    assert any("posterior" in e for e in r.errores())


def test_error_matricula_mal_formada():
    r = _resolucion_valida()
    r.inmueble.matricula_inmobiliaria = "350123456"
    assert any("350-123456" in e for e in r.errores())


def test_error_ficha_catastral_mal_formada():
    r = _resolucion_valida()
    r.inmueble.ficha_catastral = "01-09-1014"
    assert any("01-09-1014-0013-000" in e for e in r.errores())


def test_error_demolicion_sin_area():
    r = _resolucion_valida()
    r.solicitud.modalidades_ids = ["demolicion_total"]
    assert any("área a demoler" in e for e in r.errores())


def test_error_indice_propuesto_mayor_al_permitido():
    r = _resolucion_valida()
    r.inmueble.indice_ocupacion_permitido = "0,60"
    r.inmueble.indice_ocupacion_propuesto = "0,75"
    assert any("supera el permitido" in e for e in r.errores())


def test_indices_iguales_no_son_error():
    r = _resolucion_valida()
    r.inmueble.indice_construccion_permitido = "2.5"
    r.inmueble.indice_construccion_propuesto = "2.5"
    assert not any("supera el permitido" in e for e in r.errores())


def test_error_profesional_sin_matricula():
    r = _resolucion_valida()
    r.profesionales[0].matricula = ""
    assert any("matrícula profesional" in e for e in r.errores())


def test_falta_constructor_avisa_pero_no_bloquea():
    """Los profesionales son opcionales: se pueden escribir en Word."""
    r = _resolucion_valida()
    r.profesionales = [p for p in r.profesionales if p.rol != "constructor"]
    assert r.errores() == []
    assert any("constructor responsable" in a for a in r.advertencias())


def test_advertencia_titulares_mixtos():
    r = _resolucion_valida()
    r.titulares.append(Titular("JUAN PÉREZ", "M", "CC", "79000000"))
    assert any("distinto sexo" in a for a in r.advertencias())


def test_advertencia_vis_sin_estrato():
    r = _resolucion_valida()
    r.solicitud.usos_ids = ["vis_unifamiliar"]
    assert any("sin estrato" in a for a in r.advertencias())


def test_expediente_se_arma_con_el_codigo_dane():
    assert _resolucion_valida().expediente == "73001-2-26-0054"


def test_expediente_manual_tiene_prioridad():
    r = _resolucion_valida()
    r.expediente_manual = "73001-2-26-9999"
    assert r.expediente == "73001-2-26-9999"


def test_expediente_vacio_sin_consecutivo():
    r = _resolucion_valida()
    r.consecutivo_expediente = ""
    assert r.expediente == ""


def test_etiquetas_de_la_resolucion():
    e = _resolucion_valida().etiquetas()
    assert e["articulo"] == "la"
    assert e["propietario"] == "propietaria"
    assert e["fecha_letras"].startswith("quince (15) de octubre")
    assert e["area_lote_letras"] == "ciento veinte con cincuenta m2 (120,50 m2)"
    assert e["pisos"] == "dos"
    assert e["vigencia"] == "treinta y seis (36) meses"
    assert e["expediente"] == "73001-2-26-0054"
    assert e["tramite_completo"] == (
        "Licencia de CONSTRUCCIÓN en la modalidad de OBRA NUEVA")
    assert e["sentido_verbo"] == "OTORGAR"
    assert e["documentos"] == "CC 41.000.000"


def test_etiqueta_tratamiento_del_despacho():
    e = _resolucion_valida().etiquetas()
    assert e["curador_tratamiento_largo"] == (
        "La Curadora Urbana No. 2 del Municipio de Ibagué (P)")
    assert e["curador_tratamiento"] == "Curadora Urbana No. 2 de Ibagué (P)"


def test_etiquetas_de_profesionales_concuerdan():
    e = _resolucion_valida().etiquetas()
    assert e["proyectista"].startswith("la Arquitecta FERNANDA GARCÍA, "
                                       "identificada con")
    assert e["constructor"].startswith("el Arquitecto NOLBERTO GONZÁLEZ, "
                                       "identificado con")
    assert "93.383.328" in e["proyectista"]


def test_frase_solicitante_con_apoderado():
    from curaduria.resolucion import Responsable
    r = _resolucion_valida()
    r.responsable = Responsable("apoderado", "JUAN PABLO VARGAS", "M",
                                "1110454615")
    frase = r.frase_solicitante()
    assert frase.startswith("el señor JUAN PABLO VARGAS, identificado con CC. "
                            "No. 1.110.454.615")
    assert frase.endswith("en calidad de apoderado del trámite")
    assert r.errores() == []


def test_frase_solicitante_con_representante_legal():
    from curaduria.resolucion import Responsable
    r = _resolucion_valida()
    r.titulares = [Titular("CONSTRUCTORA XYZ S.A.S.", "M", "NIT",
                           "900123456-1")]
    r.responsable = Responsable("representante", "ANA RUIZ", "F", "41000001")
    frase = r.frase_solicitante()
    assert "en calidad de representante legal de CONSTRUCTORA XYZ S.A.S." in frase
    assert r.errores() == []


def test_representante_sin_persona_juridica_es_error():
    from curaduria.resolucion import Responsable
    r = _resolucion_valida()
    r.responsable = Responsable("representante", "ANA RUIZ", "F", "41000001")
    assert any("persona jurídica" in e for e in r.errores())


def test_responsable_sin_documento_es_error():
    from curaduria.resolucion import Responsable
    r = _resolucion_valida()
    r.responsable = Responsable("apoderado", "JUAN VARGAS", "M", "")
    assert any("Falta el documento de quien radica" in e for e in r.errores())


def test_frase_solicitante_sin_apoderado():
    assert _resolucion_valida().frase_solicitante().startswith("MARÍA GÓMEZ,")


def test_campos_libres_se_agregan():
    r = _resolucion_valida()
    r.extras = {"expediente_fisico": "EXP-99"}
    assert r.etiquetas()["expediente_fisico"] == "EXP-99"


def test_nombre_archivo():
    assert _resolucion_valida().nombre_archivo() == \
        "Resolucion_73001-2-26-0456_15-10-2026.docx"
    assert _resolucion_valida().nombre_archivo("hoja-datos").endswith(
        "_hoja-datos.docx")


def test_secciones_hoja_datos_cubren_los_datos_clave():
    r = _resolucion_valida()
    claves = {clave for _, campos in r.secciones_hoja_datos()
              for _, clave in campos}
    for clave in ("resolucion_numero", "expediente", "matricula",
                  "ficha_catastral", "nombres", "proyectista_nombre",
                  "idu_certificado", "sustanciador"):
        assert clave in claves, clave


def test_subtitulo_hoja_datos():
    assert "Expediente 73001-2-26-0054" in \
        _resolucion_valida().subtitulo_hoja_datos()


def test_formato_documento():
    from curaduria.gramatica import formato_documento
    assert formato_documento("79351573") == "79.351.573"
    assert formato_documento("900123456-1") == "900.123.456-1"
    assert formato_documento("") == ""
    assert formato_documento("sin numero") == "sin numero"


def test_plantilla_maestra_se_rellena_completa(tmp_path):
    """La plantilla de ejemplo no debe dejar etiquetas de datos sin resolver."""
    import crear_plantilla_ejemplo

    ruta = tmp_path / "p.docx"
    crear_plantilla_ejemplo.construir().save(ruta)
    doc = plantilla.abrir(ruta)
    etiquetas = _resolucion_valida().etiquetas()
    faltantes = plantilla.rellenar(doc, etiquetas)
    # Las anclas de sección no son datos: las llenan los bloques.
    assert faltantes == {"encabezado", "considerandos", "resuelve", "cierre"}
    texto = "\n".join(p.text for p in plantilla.recorrer_parrafos(doc))
    assert "73001-2-26-0054" in texto
    assert "Licencia de CONSTRUCCIÓN en la modalidad de OBRA NUEVA" in texto
    assert "Curaduría Urbana No. 2 de Ibagué" in texto


# --- Avance del formulario (Nivel 1) --------------------------------------

def test_el_formulario_se_encoge_segun_el_tramite():
    from curaduria.catalogo import Solicitud

    construccion = _resolucion_valida()
    construccion.solicitud = Solicitud("construccion", ["obra_nueva"],
                                       ["vivienda_unifamiliar"])
    concepto = Resolucion(solicitud=Solicitud("concepto_uso"))
    assert len(concepto.estado_campos()) < len(construccion.estado_campos())
    nombres = {c.nombre for c in concepto.estado_campos()}
    assert "Pisos y cubierta" not in nombres


def test_la_demolicion_pide_el_area_a_demoler():
    from curaduria.catalogo import Solicitud

    r = _resolucion_valida()
    r.solicitud = Solicitud("construccion", ["demolicion_total"])
    assert "Área a demoler" in {c.nombre for c in r.estado_campos()}
    r.solicitud = Solicitud("construccion", ["cerramiento"])
    assert "Área a demoler" not in {c.nombre for c in r.estado_campos()}


def test_los_indices_solo_en_obra_nueva_o_ampliacion():
    from curaduria.catalogo import Solicitud

    r = _resolucion_valida()
    r.solicitud = Solicitud("construccion", ["obra_nueva"])
    assert any("Índices" in c.nombre for c in r.estado_campos())
    r.solicitud = Solicitud("construccion", ["demolicion_total"])
    assert not any("Índices" in c.nombre for c in r.estado_campos())


def test_el_estrato_aparece_solo_en_vis_o_vip():
    from curaduria.catalogo import Solicitud

    r = _resolucion_valida()
    assert not any("Estrato" in c.nombre for c in r.estado_campos())
    r.solicitud = Solicitud("construccion", ["obra_nueva"], ["vip"])
    assert any("Estrato (proyecto VIP)" in c.nombre for c in r.estado_campos())


def test_avance_cuenta_los_campos_listos():
    r = _resolucion_valida()
    listos, total = r.avance
    assert 0 < listos <= total
    r.numero = ""
    menos, _ = r.avance
    assert menos < listos


def test_el_detalle_de_titulares_dice_la_concordancia():
    r = _resolucion_valida()
    campo = next(c for c in r.estado_campos() if c.nombre == "Titulares")
    assert campo.detalle == "1: la propietaria"
    r.titulares.append(Titular("JUAN PÉREZ", "M", "CC", "79000000"))
    campo = next(c for c in r.estado_campos() if c.nombre == "Titulares")
    assert campo.detalle == "2: los propietarios"


def test_quien_radica_el_propio_titular():
    r = _resolucion_valida()
    campo = next(c for c in r.estado_campos() if c.nombre == "Quién radica")
    assert campo.listo and campo.detalle == "el propio titular"


def test_quien_radica_apoderado_incompleto():
    from curaduria.resolucion import Responsable

    r = _resolucion_valida()
    r.responsable = Responsable("apoderado")
    campo = next(c for c in r.estado_campos() if c.nombre == "Quién radica")
    assert not campo.listo
    assert "falta el apoderado" in campo.detalle


def test_la_ficha_catastral_no_es_critica():
    r = _resolucion_valida()
    campo = next(c for c in r.estado_campos()
                 if c.nombre == "Ficha catastral")
    assert campo.critico is False
