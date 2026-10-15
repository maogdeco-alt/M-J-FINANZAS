"""Armado del documento: inserción de bloques y hoja de datos."""

from datetime import date

from docx import Document

from curaduria import armado, plantilla
from curaduria.bloques import Biblioteca, Bloque
from curaduria.catalogo import Solicitud
from curaduria.gramatica import Titular
from curaduria.profesionales import Profesional
from curaduria.resolucion import Despacho, Inmueble, Resolucion


def _resolucion():
    return Resolucion(
        numero="73001-2-26-0456", fecha=date(2026, 10, 15), radicado="2600789",
        fecha_radicado=date(2026, 3, 10), consecutivo_expediente="54",
        solicitud=Solicitud("construccion", ["demolicion_total", "obra_nueva"],
                            ["vivienda_unifamiliar"]),
        despacho=Despacho(curador="SANDRA PARDO", sustanciador="GLORIA LOZANO",
                          sustanciador_iniciales="gilt /S"),
        titulares=[Titular("MARÍA GÓMEZ", "F", "CC", "41000000"),
                   Titular("JUAN PÉREZ", "M", "CC", "79000000")],
        inmueble=Inmueble(direccion="CL 10 # 5-20",
                          matricula_inmobiliaria="350-123456",
                          ficha_catastral="01-09-1014-0013-000",
                          area_lote=120.5, area_construida=85.0,
                          area_demoler=40.0, pisos=2, estrato="2"),
        profesionales=[
            Profesional("proyectista", "FERNANDA GARCÍA", "F", "93383328", "25700-1"),
            Profesional("constructor", "NOLBERTO GONZÁLEZ", "M", "93377891", "A164-1"),
        ])


# --- Inserción en la plantilla --------------------------------------------

def _doc_con_ancla(ancla="considerandos"):
    doc = Document()
    doc.add_paragraph("ANTES")
    parrafo = doc.add_paragraph()
    run = parrafo.add_run(f"{{{{{ancla}}}}}")
    run.bold = True
    doc.add_paragraph("DESPUÉS")
    return doc


def test_insertar_textos_reemplaza_el_ancla():
    doc = _doc_con_ancla()
    assert armado.insertar_textos(doc, "considerandos", ["uno", "dos", "tres"])
    textos = [p.text for p in doc.paragraphs]
    assert textos == ["ANTES", "uno", "dos", "tres", "DESPUÉS"]


def test_insertar_textos_hereda_el_formato_del_ancla():
    doc = _doc_con_ancla()
    armado.insertar_textos(doc, "considerandos", ["uno"])
    insertado = [p for p in doc.paragraphs if p.text == "uno"][0]
    assert insertado.runs[0].bold is True


def test_insertar_textos_sin_ancla():
    doc = Document()
    doc.add_paragraph("solo texto")
    assert armado.insertar_textos(doc, "considerandos", ["x"]) is False
    assert [p.text for p in doc.paragraphs] == ["solo texto"]


def test_insertar_textos_en_ancla_con_espacios():
    doc = Document()
    doc.add_paragraph("{{ considerandos }}")
    assert armado.insertar_textos(doc, "considerandos", ["uno"])
    assert [p.text for p in doc.paragraphs] == ["uno"]


# --- Hoja de datos ---------------------------------------------------------

def test_filas_hoja_datos():
    filas = armado.filas_hoja_datos(
        {"a": "1", "b": ""},
        [("Sección", [("Campo A", "a"), ("Campo B", "b"),
                      ("Campo C", "inexistente")])])
    assert filas == [("Sección", "Campo A", "1"), ("Sección", "Campo B", ""),
                     ("Sección", "Campo C", "")]


def test_hoja_datos_escribe_la_tabla():
    filas = [("Acto", "Resolución No.", "0456"), ("Acto", "Fecha", ""),
             ("Predio", "Matrícula", "350-123456")]
    doc = armado.documento_hoja_datos(filas, "HOJA DE DATOS", "subtítulo")
    tabla = doc.tables[0]
    # dos filas de sección + tres de campos
    assert len(tabla.rows) == 5
    texto = "\n".join(c.text for fila in tabla.rows for c in fila.cells)
    assert "ACTO" in texto and "PREDIO" in texto
    assert "350-123456" in texto


def test_hoja_datos_marca_los_pendientes_con_raya():
    doc = armado.documento_hoja_datos(
        [("Acto", "Fecha", "")], "HOJA", "")
    fila = doc.tables[0].rows[1]
    assert "____" in fila.cells[1].text


def test_hoja_datos_de_una_resolucion_real():
    resolucion = _resolucion()
    filas = armado.filas_hoja_datos(resolucion.etiquetas(),
                                    resolucion.secciones_hoja_datos())
    doc = armado.documento_hoja_datos(filas, "HOJA DE DATOS",
                                      resolucion.subtitulo_hoja_datos())
    texto = "\n".join(c.text for t in doc.tables
                      for f in t.rows for c in f.cells)
    assert "73001-2-26-0054" in texto
    assert "MARÍA GÓMEZ y JUAN PÉREZ" in texto
    assert "01-09-1014-0013-000" in texto
    assert "FERNANDA GARCÍA" in texto


# --- Armado completo -------------------------------------------------------

def test_armar_desde_bloques_sin_plantilla():
    libreria = Biblioteca()
    resolucion = _resolucion()
    etiquetas = resolucion.etiquetas()
    bloques = libreria.obligatorios("construccion")
    doc, faltantes = armado.armar_desde_bloques(bloques, etiquetas)
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert "RESOLUCIÓN No. 73001-2-26-0456" in texto
    assert "RESUELVE" in texto
    assert "CONSIDERANDO" in texto
    assert "MARÍA GÓMEZ, identificada con cédula de ciudadanía No. 41.000.000" \
        in texto
    assert "JUAN PÉREZ, identificado con cédula" in texto
    assert "en calidad de propietarios" in texto   # plural por dos titulares
    assert "73001-2-26-0054" in texto
    assert texto.index("CONSIDERANDO") < texto.index("RESUELVE")


def test_armar_desde_bloques_incluye_la_hoja_de_datos():
    libreria = Biblioteca()
    resolucion = _resolucion()
    etiquetas = resolucion.etiquetas()
    filas = armado.filas_hoja_datos(etiquetas, resolucion.secciones_hoja_datos())
    doc, _ = armado.armar_desde_bloques(
        libreria.seleccion(["preambulo", "solicitud"]), etiquetas, filas)
    assert doc.tables, "la hoja de datos debe quedar como tabla"


def test_armar_sobre_plantilla_inserta_en_las_anclas():
    import crear_plantilla_ejemplo

    libreria = Biblioteca()
    resolucion = _resolucion()
    etiquetas = resolucion.etiquetas()
    doc = crear_plantilla_ejemplo.construir()
    bloques = libreria.seleccion(
        ["solicitud", "citacion_vecinos", "art_plazo", "art_recursos"])
    informe = armado.armar_sobre_plantilla(doc, bloques, etiquetas)

    assert informe["ancla_considerandos"] is True
    assert informe["ancla_resuelve"] is True
    assert informe["bloques_por_seccion"]["considerando"] == 2
    assert informe["bloques_por_seccion"]["resuelve"] == 2
    assert informe["anclas_faltantes"] == []

    texto = "\n".join(p.text for p in plantilla.recorrer_parrafos(doc))
    assert "{{considerandos}}" not in texto and "{{resuelve}}" not in texto
    assert "Licencia de CONSTRUCCIÓN en las modalidades de DEMOLICIÓN TOTAL y "\
           "OBRA NUEVA" in texto
    assert "350-123456" in texto
    assert "treinta y seis (36) meses" in texto
    # El orden del catálogo se respeta: la solicitud antes de la citación.
    assert texto.index("solicita Licencia") < texto.index("Citación a Vecinos")


def test_armar_sobre_plantilla_reporta_faltantes_de_los_bloques():
    import crear_plantilla_ejemplo

    libreria = Biblioteca()
    doc = crear_plantilla_ejemplo.construir()
    resolucion = _resolucion()
    informe = armado.armar_sobre_plantilla(
        doc, libreria.seleccion(["reparto"]), resolucion.etiquetas())
    assert "reparto_numero" in informe["faltantes"]


def test_armar_sobre_plantilla_sin_bloques_solo_rellena():
    import crear_plantilla_ejemplo

    doc = crear_plantilla_ejemplo.construir()
    informe = armado.armar_sobre_plantilla(doc, [], _resolucion().etiquetas())
    assert informe["ancla_considerandos"] is False
    texto = "\n".join(p.text for p in plantilla.recorrer_parrafos(doc))
    assert "{{considerandos}}" in texto  # el ancla queda visible para rellenar


def test_bloque_con_texto_largo_mantiene_sus_parrafos():
    bloque = Bloque(id="x", titulo="X",
                    texto="Primero.\n\nSegundo.\n\nTercero.")
    doc = Document()
    armado.agregar_textos_al_final(doc, bloque.render({}))
    assert [p.text for p in doc.paragraphs] == ["Primero.", "Segundo.", "Tercero."]


# --- Anclas por sección ----------------------------------------------------

def _plantilla_con(anclas):
    doc = Document()
    for ancla in anclas:
        doc.add_paragraph(f"{{{{{ancla}}}}}")
    return doc


def test_cada_seccion_va_a_su_ancla():
    libreria = Biblioteca()
    doc = _plantilla_con(["encabezado", "considerandos", "resuelve", "cierre"])
    bloques = libreria.seleccion(["preambulo", "solicitud", "art_plazo", "cierre"])
    informe = armado.armar_sobre_plantilla(doc, bloques, _resolucion().etiquetas())
    assert informe["anclas_faltantes"] == []
    assert informe["secciones_desviadas"] == []
    textos = [p.text for p in doc.paragraphs if p.text.strip()]
    posicion = {nombre: i for i, t in enumerate(textos) for nombre, marca in (
        ("preambulo", "en ejercicio de las funciones"),
        ("solicitud", "solicita Licencia"),
        ("plazo", "tendrá una vigencia"),
        ("cierre", "NOTIFÍQUESE")) if marca in t}
    assert (posicion["preambulo"] < posicion["solicitud"]
            < posicion["plazo"] < posicion["cierre"])


def test_sin_ancla_de_encabezado_el_preambulo_encabeza_los_considerandos():
    libreria = Biblioteca()
    doc = _plantilla_con(["considerandos", "resuelve"])
    bloques = libreria.seleccion(["preambulo", "solicitud"])
    informe = armado.armar_sobre_plantilla(doc, bloques, _resolucion().etiquetas())
    assert informe["secciones_desviadas"] == ["encabezado"]
    assert "encabezado" in informe["anclas_faltantes"]
    textos = [p.text for p in doc.paragraphs if p.text.strip()]
    assert "en ejercicio de las funciones" in textos[0]
    assert "solicita Licencia" in " ".join(textos)


def test_sin_ancla_de_cierre_el_cierre_va_con_la_parte_resolutiva():
    libreria = Biblioteca()
    doc = _plantilla_con(["considerandos", "resuelve"])
    bloques = libreria.seleccion(["art_plazo", "cierre", "pie_archivo"])
    informe = armado.armar_sobre_plantilla(doc, bloques, _resolucion().etiquetas())
    assert informe["secciones_desviadas"] == ["cierre"]
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert "NOTIFÍQUESE" in texto
    assert texto.index("tendrá una vigencia") < texto.index("NOTIFÍQUESE")


def test_sin_ninguna_ancla_no_se_pierde_el_relleno_de_datos():
    doc = Document()
    doc.add_paragraph("Resolución No. {{resolucion_numero}}")
    informe = armado.armar_sobre_plantilla(
        doc, Biblioteca().seleccion(["art_plazo"]), _resolucion().etiquetas())
    assert informe["ancla_resuelve"] is False
    assert doc.paragraphs[0].text == "Resolución No. 73001-2-26-0456"


def test_hoja_de_datos_en_el_ancla():
    doc = _plantilla_con(["hoja_datos", "considerandos"])
    doc.add_paragraph("FIN")
    informe = armado.armar_sobre_plantilla(
        doc, [], _resolucion().etiquetas(),
        [("Acto", "Resolución No.", "0456")])
    assert informe["hoja_datos"] == "ancla"
    assert doc.tables
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert "{{hoja_datos}}" not in texto


def test_hoja_de_datos_sin_ancla_queda_al_inicio():
    doc = _plantilla_con(["considerandos"])
    informe = armado.armar_sobre_plantilla(
        doc, [], _resolucion().etiquetas(),
        [("Acto", "Resolución No.", "0456")])
    assert informe["hoja_datos"] == "inicio"
    assert doc.paragraphs[0].text == "HOJA DE DATOS DEL EXPEDIENTE"


def test_plantilla_de_ejemplo_no_duplica_el_preambulo_ni_el_cierre():
    import crear_plantilla_ejemplo

    libreria = Biblioteca()
    resolucion = _resolucion()
    doc = crear_plantilla_ejemplo.construir()
    bloques = libreria.obligatorios("construccion")
    armado.armar_sobre_plantilla(doc, bloques, resolucion.etiquetas())
    texto = "\n".join(p.text for p in plantilla.recorrer_parrafos(doc))
    assert texto.count("en ejercicio de las funciones otorgadas") == 1
    assert texto.count("se notifica personalmente") == 1
    assert texto.count("NOTIFÍQUESE") == 1
