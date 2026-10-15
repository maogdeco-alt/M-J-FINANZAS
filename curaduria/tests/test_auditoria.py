"""Revisión de plantillas: etiquetas, anclas y espacios en blanco."""

from docx import Document

from curaduria.auditoria import auditar, etiquetas_conocidas


def _doc(*textos):
    documento = Document()
    for texto in textos:
        documento.add_paragraph(texto)
    return documento


def test_etiquetas_conocidas_incluye_datos_anclas_y_bloques():
    conocidas = etiquetas_conocidas()
    for clave in ("matricula", "nombres", "fecha_letras", "proyectista",
                  "considerandos", "resuelve", "hoja_datos",
                  "solicitante_frase", "reparto_numero"):
        assert clave in conocidas, clave


def test_detecta_las_etiquetas_usadas():
    informe = auditar(_doc("Predio {{matricula}} de {{municipio}}."))
    assert informe.etiquetas == ["matricula", "municipio"]
    assert informe.etiquetas_desconocidas == []
    assert informe.sirve


def test_detecta_etiqueta_mal_escrita():
    informe = auditar(_doc("Predio {{matrikula}}."))
    assert informe.etiquetas_desconocidas == ["matrikula"]
    assert not informe.sirve


def test_detecta_las_anclas():
    informe = auditar(_doc("{{considerandos}}", "{{resuelve}}"))
    assert informe.anclas == ["considerandos", "resuelve"]
    assert "cierre" in informe.anclas_faltantes


def test_detecta_rayas():
    informe = auditar(_doc("El señor ____________ presenta."))
    assert len(informe.blancos) == 1
    assert informe.blancos[0].tipo == "raya"


def test_detecta_equis():
    informe = auditar(_doc("Resolución No. 0XXX de 2026."))
    assert [h.texto for h in informe.blancos] == ["0XXX"]
    assert informe.blancos[0].tipo == "equis_pegada"


def test_no_marca_lo_que_ya_es_etiqueta():
    informe = auditar(_doc("Matrícula {{matricula}} del predio."))
    assert informe.blancos == []
    assert informe.lista


def test_sugiere_la_etiqueta_segun_el_contexto():
    informe = auditar(_doc(
        "identificado con Matrícula Inmobiliaria No. ________ y "
        "Ficha Catastral No. ________"))
    assert [h.sugerencia for h in informe.blancos] == ["matricula",
                                                       "ficha_catastral"]


def test_sugiere_fecha_cuando_el_blanco_es_parte_de_una_fecha():
    informe = auditar(_doc("entregada personalmente el XX de ______ de 2026"))
    assert {h.sugerencia for h in informe.blancos} == {"fecha_larga"}


def test_los_blancos_salen_en_el_orden_del_texto():
    informe = auditar(_doc("No. 0XXX del XX de ____ de 2026"))
    posiciones = [h.posicion for h in informe.blancos]
    assert posiciones == sorted(posiciones)


def test_un_mismo_tramo_no_se_reporta_dos_veces():
    informe = auditar(_doc("XXXX"))
    assert len(informe.blancos) == 1


def test_revisa_tablas_y_encabezados():
    documento = Document()
    documento.add_table(rows=1, cols=1).cell(0, 0).text = "{{matricula}} ____"
    documento.sections[0].header.paragraphs[0].text = "{{curaduria}}"
    informe = auditar(documento)
    assert "matricula" in informe.etiquetas
    assert "curaduria" in informe.etiquetas
    assert len(informe.blancos) == 1


def test_resumen_legible():
    informe = auditar(_doc("{{matricula}} y ____"))
    assert "1 etiquetas" in informe.resumen()
    assert "1 espacios en blanco" in informe.resumen()


def test_plantilla_de_ejemplo_esta_lista():
    import crear_plantilla_ejemplo

    informe = auditar(crear_plantilla_ejemplo.construir())
    assert informe.lista, informe.etiquetas_desconocidas or informe.blancos
    assert set(informe.anclas) >= {"encabezado", "considerandos", "resuelve",
                                   "cierre"}
