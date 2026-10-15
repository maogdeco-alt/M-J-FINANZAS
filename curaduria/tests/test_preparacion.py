"""Conversión automática de una plantilla real en plantilla de la app."""

from docx import Document

from curaduria.auditoria import convertir, convertir_texto, parece_fecha
from curaduria.plantilla import recorrer_parrafos
from curaduria.preparacion import poner_anclas, preparar, reemplazar_literales


def _doc(*textos):
    documento = Document()
    for texto in textos:
        documento.add_paragraph(texto)
    return documento


def _texto(doc):
    return "\n".join(p.text for p in recorrer_parrafos(doc) if p.text.strip())


# --- Conversión de espacios en blanco -------------------------------------

def test_convierte_una_matricula_en_blanco():
    doc = _doc("identificado con Matrícula Inmobiliaria No. ________")
    cambios, pendientes = convertir(doc)
    assert _texto(doc).endswith("No. {{matricula}}")
    assert [c.etiqueta for c in cambios] == ["matricula"]
    assert pendientes == []


def test_una_fecha_partida_en_tres_blancos_queda_en_una_etiqueta():
    doc = _doc("entregada personalmente el XX de ______ de 2026, conforme")
    convertir(doc)
    assert _texto(doc) == "entregada personalmente el {{fecha_larga}}, conforme"


def test_fecha_con_anio_incompleto():
    doc = _doc("quedando en legal y debida forma el XX de ____ de 202X,")
    convertir(doc)
    assert _texto(doc) == "quedando en legal y debida forma el {{fecha_larga}},"


def test_el_numero_de_acto_absorbe_su_prefijo():
    """Si el prefijo se deja, el documento sale con el número repetido."""
    doc = _doc("mediante Auto No. 73001-2-26-0XXX de la misma fecha")
    convertir(doc)
    assert _texto(doc) == "mediante Auto No. {{auto_numero}} de la misma fecha"


def test_no_inventa_etiqueta_cuando_no_hay_pista():
    doc = _doc("publicado en el Boletín de Curadores urbanos No. 26-XXXX")
    cambios, pendientes = convertir(doc)
    assert cambios == []
    assert len(pendientes) == 1
    assert "{{" not in _texto(doc)


def test_lo_que_se_convierte_no_queda_como_pendiente():
    doc = _doc("el XX de ______ de 2026")
    cambios, pendientes = convertir(doc)
    assert len(cambios) == 1 and pendientes == []


def test_sin_aplicar_no_toca_el_documento():
    doc = _doc("Matrícula Inmobiliaria No. ________")
    antes = _texto(doc)
    cambios, _ = convertir(doc, aplicar=False)
    assert cambios and _texto(doc) == antes


def test_conserva_el_formato_al_convertir():
    documento = Document()
    parrafo = documento.add_paragraph()
    run = parrafo.add_run("Matrícula Inmobiliaria No. ________")
    run.bold = True
    convertir(documento)
    assert documento.paragraphs[0].runs[0].bold is True
    assert "{{matricula}}" in documento.paragraphs[0].text


def test_convierte_en_tablas_y_encabezados():
    documento = Document()
    celda = documento.add_table(rows=1, cols=1).cell(0, 0)
    celda.paragraphs[0].add_run("Ficha Catastral No. ________")
    documento.sections[0].header.paragraphs[0].add_run("RESOLUCION No. 0XXX")
    convertir(documento)
    assert "{{ficha_catastral}}" in celda.text
    assert "{{resolucion_numero}}" in \
        documento.sections[0].header.paragraphs[0].text


# --- Forma de fecha --------------------------------------------------------

def test_parece_fecha():
    assert parece_fecha("XX de ______ de 2026")
    assert parece_fecha("___ de _____ de 20XX")
    assert not parece_fecha("25000XXXX de ___________")
    assert not parece_fecha("")


def test_convertir_texto_solo_fechas_deja_lo_demas():
    texto = ("mediante Oficio No. 25000XXXX de ___________ y el "
             "XX de ______ de 2026")
    nuevo, cambios = convertir_texto(texto, solo_fechas=True)
    assert [c.etiqueta for c in cambios] == ["fecha_larga"]
    assert "25000XXXX" in nuevo
    assert nuevo.endswith("y el {{fecha_larga}}")


# --- Datos del despacho escritos literalmente -----------------------------

def test_reemplaza_el_nombre_de_la_curadora():
    doc = _doc("SANDRA DEL PILAR PARDO SUAREZ", "Curadora Urbana de Ibagué")
    cambios = reemplazar_literales(
        doc, {"curador": "SANDRA DEL PILAR PARDO SUAREZ",
              "municipio": "Ibagué"})
    assert "{{curador}}" in _texto(doc)
    assert "{{municipio}}" in _texto(doc)
    assert len(cambios) == 2


def test_no_toca_lo_que_ya_es_etiqueta():
    doc = _doc("Expedida en {{municipio}} por Ibagué")
    reemplazar_literales(doc, {"municipio": "Ibagué"})
    assert _texto(doc).count("{{municipio}}") == 2


def test_ignora_valores_muy_cortos():
    doc = _doc("La curaduría No. 2 de Ibagué")
    cambios = reemplazar_literales(doc, {"curaduria_numero": "2"})
    assert cambios == []


# --- Anclas de sección -----------------------------------------------------

def _plantilla_real_simulada():
    return _doc(
        "RESOLUCIÓN No. 0456",
        "La Curadora Urbana No. 2 en ejercicio de las funciones otorgadas",
        "por la Ley 388 de 1997 y el Decreto 1077 de 2015 y,",
        "CONSIDERANDO",
        "Que el señor X solicita licencia.",
        "VIS",
        "Que allega Declaración de Vivienda de Interés Social.",
        "Que se resuelve dicha intervención y por tanto será tenida en cuenta.",
        "RESUELVE",
        "Artículo 1.- OTORGAR la licencia.",
        "Artículo 2.- PLAZO: treinta y seis meses.",
        "NOTIFÍQUESE, PUBLÍQUESE Y CÚMPLASE",
        "SANDRA PARDO",
    )


def test_pone_las_cuatro_anclas():
    doc = _plantilla_real_simulada()
    anclas, recortados, nota = poner_anclas(doc)
    assert anclas == ["encabezado", "considerandos", "resuelve", "cierre"]
    assert nota == ""
    assert recortados > 0
    assert _texto(doc).split("\n") == [
        "RESOLUCIÓN No. 0456", "{{encabezado}}", "CONSIDERANDO",
        "{{considerandos}}", "RESUELVE", "{{resuelve}}", "{{cierre}}"]


def test_el_preambulo_partido_en_dos_queda_en_una_sola_ancla():
    doc = _plantilla_real_simulada()
    poner_anclas(doc)
    assert _texto(doc).count("{{encabezado}}") == 1
    assert "Decreto 1077 de 2015 y," not in _texto(doc)


def test_resuelve_dentro_de_una_frase_no_corta_el_documento():
    """"se resuelve dicha intervención" es un considerando, no el título."""
    doc = _plantilla_real_simulada()
    poner_anclas(doc)
    texto = _texto(doc)
    assert texto.index("{{considerandos}}") < texto.index("RESUELVE")
    assert "intervención" not in texto


def test_sin_el_titulo_considerando_no_pone_anclas_y_avisa():
    doc = _doc("Un documento cualquiera", "sin estructura de resolución")
    anclas, recortados, nota = poner_anclas(doc)
    assert anclas == [] and recortados == 0
    assert "CONSIDERANDO" in nota


def test_sin_el_titulo_resuelve_avisa():
    doc = _doc("CONSIDERANDO", "Que algo.")
    anclas, _, nota = poner_anclas(doc)
    assert anclas == [] and "RESUELVE" in nota


# --- Todo junto -------------------------------------------------------------

def test_preparar_hace_las_tres_cosas():
    doc = _plantilla_real_simulada()
    doc.add_paragraph("Matrícula Inmobiliaria No. ________ de Ibagué")
    informe = preparar(doc, {"municipio": "Ibagué"})
    assert informe.etiquetas_puestas
    assert informe.literales
    assert informe.anclas == ["encabezado", "considerandos", "resuelve",
                              "cierre"]
    assert "etiquetas" in informe.resumen()


def test_preparar_sin_anclas():
    doc = _plantilla_real_simulada()
    informe = preparar(doc, None, con_anclas=False)
    assert informe.anclas == []
    assert "CONSIDERANDO" in _texto(doc)


def test_la_plantilla_preparada_queda_lista_para_la_app():
    """Después de preparar, la auditoría no debe encontrar nada pendiente."""
    from curaduria.auditoria import auditar

    doc = _plantilla_real_simulada()
    doc.add_paragraph("Ficha Catastral No. ________ del XX de ____ de 2026")
    preparar(doc, {"municipio": "Ibagué"})
    informe = auditar(doc)
    assert informe.sirve
    assert informe.blancos == []
    assert set(informe.anclas) >= {"encabezado", "considerandos", "resuelve",
                                   "cierre"}
