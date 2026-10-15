"""Pruebas de la interfaz: que los campos guarden y que los botones funcionen.

Son las dos quejas que importan en el uso diario. `AppTest` corre la app de
verdad, sin navegador, así que un campo que se borre o un botón que no haga
nada se detecta aquí y no en la oficina.
"""

from pathlib import Path

import pytest
from streamlit.testing.v1 import AppTest

RAIZ = Path(__file__).resolve().parent.parent
APP = str(RAIZ / "app.py")


@pytest.fixture
def app(tmp_path, monkeypatch):
    """La app con una configuración de prueba, ya completa."""
    import curaduria.configuracion as configuracion

    archivo = tmp_path / "configuracion.json"
    monkeypatch.setattr(configuracion, "ARCHIVO", archivo)
    configuracion.Configuracion(
        curador="SANDRA DEL PILAR PARDO SUAREZ",
        sustanciador="GLORIA INES LOZANO TRUJILLO",
        sustanciador_iniciales="gilt /S").guardar(archivo)

    prueba = AppTest.from_file(APP, default_timeout=180)
    prueba.run()
    assert not prueba.exception, [str(e) for e in prueba.exception]
    return prueba


def campo(app, etiqueta):
    return next((c for c in app.text_input if c.label == etiqueta), None)


def numero(app, etiqueta):
    return next((n for n in app.number_input if n.label == etiqueta), None)


def boton(app, etiqueta):
    return next((b for b in app.button if b.label == etiqueta), None)


def llenar_minimo(app):
    """Los datos que exige una licencia de construcción de obra nueva."""
    app.multiselect[0].set_value(["OBRA NUEVA"]).run()
    app.multiselect[1].set_value(["VIVIENDA UNIFAMILIAR"]).run()
    campo(app, "Resolución No.").set_value("73001-2-26-0456").run()
    campo(app, "Radicación No.").set_value("2600789").run()
    campo(app, "Consecutivo del expediente").set_value("54").run()
    campo(app, "Matrícula inmobiliaria").set_value("350250605").run()
    campo(app, "Dirección del predio").set_value("cra 2 sur no 28-57").run()
    app.session_state[ "_datos_guardados"]["titulares"] = \
        app.session_state["_datos_guardados"]["titulares"].assign(
            nombre=["MARIA LELI FIRIGUA"], numero_doc=["28514217"],
            sexo=["Femenino"])
    app.run()
    numero(app, "Área del lote (m²)").set_value(150.0).run()
    numero(app, "Área construida (m²)").set_value(245.11).run()
    numero(app, "Pisos").set_value(2).run()
    return app


# --- La app abre ------------------------------------------------------------

def test_la_app_abre_sin_errores(app):
    assert len(app.tabs) == 4
    assert not app.exception


def test_el_despacho_configurado_se_muestra(app):
    assert any("SANDRA DEL PILAR PARDO SUAREZ" in m.value
               for m in app.markdown)


# --- Los campos guardan lo escrito -----------------------------------------

def test_un_campo_conserva_lo_escrito(app):
    campo(app, "Matrícula inmobiliaria").set_value("350250605").run()
    assert campo(app, "Matrícula inmobiliaria").value == "350-250605"
    # otra interacción cualquiera no debe borrarlo
    campo(app, "Radicación No.").set_value("2600789").run()
    assert campo(app, "Matrícula inmobiliaria").value == "350-250605"


def test_un_campo_que_se_oculta_y_vuelve_no_pierde_su_valor(app):
    """El formulario se encoge según la modalidad: lo escrito debe sobrevivir."""
    app.multiselect[0].set_value(["DEMOLICIÓN TOTAL", "OBRA NUEVA"]).run()
    numero(app, "Área a demoler (m²)").set_value(113.5).run()

    app.multiselect[0].set_value(["OBRA NUEVA"]).run()
    assert numero(app, "Área a demoler (m²)") is None   # se ocultó

    app.multiselect[0].set_value(["DEMOLICIÓN TOTAL", "OBRA NUEVA"]).run()
    assert numero(app, "Área a demoler (m²)").value == 113.5


def test_cambiar_de_tramite_no_borra_los_datos_comunes(app):
    campo(app, "Resolución No.").set_value("73001-2-26-0456").run()
    app.selectbox[0].set_value("reconocimiento").run()   # cambia el grupo
    assert campo(app, "Resolución No.").value == "73001-2-26-0456"


def test_la_normalizacion_corrige_al_escribir(app):
    campo(app, "Ficha catastral").set_value("010910140013000").run()
    assert campo(app, "Ficha catastral").value == "01-09-1014-0013-000"
    campo(app, "Dirección del predio").set_value("cra 2 sur no 28-57").run()
    assert campo(app, "Dirección del predio").value == "Carrera 2 Sur No. 28-57"


# --- Los botones funcionan --------------------------------------------------

def test_boton_sugerir_sexo(app):
    almacen = app.session_state["_datos_guardados"]
    almacen["titulares"] = almacen["titulares"].assign(
        nombre=["MARIA LELI FIRIGUA"])
    app.run()
    boton(app, "Sugerir el sexo por el nombre").click().run()
    assert app.session_state["_datos_guardados"]["titulares"].iloc[0]["sexo"] \
        == "Femenino"


def test_botones_de_seleccion_de_parrafos(app):
    boton(app, "Marcar todos").click().run()
    todos = len(app.session_state.bloques_elegidos)
    boton(app, "Solo obligatorios").click().run()
    obligatorios = len(app.session_state.bloques_elegidos)
    assert 0 < obligatorios < todos


def test_boton_limpiar_el_formulario(app):
    campo(app, "Resolución No.").set_value("73001-2-26-0456").run()
    boton(app, "Limpiar el formulario").click().run()
    assert campo(app, "Resolución No.").value == ""


def test_el_boton_de_generar_esta_bloqueado_si_faltan_datos(app):
    generar = boton(app, "GENERAR LA RESOLUCIÓN EN WORD")
    assert generar.disabled


def test_generar_la_resolucion(app):
    llenar_minimo(app)
    assert not [e.value for e in app.error], [e.value for e in app.error]
    boton(app, "GENERAR LA RESOLUCIÓN EN WORD").click().run()
    assert not app.exception
    assert app.session_state.documento
    assert app.session_state.documento[:2] == b"PK"   # es un .docx
    assert any("Descargar" in d.label.title() or "DESCARGAR" in d.label
               for d in app.download_button)


def test_generar_la_hoja_de_datos(app):
    llenar_minimo(app)
    boton(app, "Preparar la hoja de datos").click().run()
    assert app.session_state.hoja[:2] == b"PK"


def test_guardar_la_configuracion(app, tmp_path):
    import json

    for entrada in app.text_input:
        if entrada.label == "Municipio":
            entrada.set_value("Espinal")
    next(b for b in app.button
         if b.label == "GUARDAR LA CONFIGURACIÓN").click().run()
    assert app.session_state.configuracion.municipio == "Espinal"
    guardado = json.loads((tmp_path / "configuracion.json").read_text("utf-8"))
    assert guardado["municipio"] == "Espinal"


def test_el_avance_sube_al_llenar(app):
    antes = app.progress[0].value
    llenar_minimo(app)
    assert app.progress[0].value > antes


# --- Marca ------------------------------------------------------------------

def test_el_logo_se_muestra(app):
    assert len(app.image) >= 2, "el logo va en el membrete y en la barra lateral"


def test_el_documento_generado_lleva_membrete(app):
    import io

    from docx import Document
    from docx.oxml.ns import qn

    llenar_minimo(app)
    boton(app, "GENERAR LA RESOLUCIÓN EN WORD").click().run()
    documento = Document(io.BytesIO(app.session_state.documento))
    encabezado = documento.sections[0].header
    tiene_imagen = any(
        run.element.findall(qn("a:blip")) or run.element.findall(qn("w:drawing"))
        for parrafo in encabezado.paragraphs for run in parrafo.runs)
    tiene_texto = any(p.text.strip() for p in encabezado.paragraphs)
    assert tiene_imagen or tiene_texto
