"""Relleno de plantillas de Word conservando el formato original.

python-docx parte el texto de un párrafo en *runs* según el formato, así que
una etiqueta como {{nombre}} puede quedar repartida en varios runs (sobre todo
si alguien la escribió con el corrector encima). El algoritmo de aquí reconstruye
el texto del párrafo, hace el reemplazo y lo reescribe conservando el formato
del run donde arranca cada etiqueta.
"""

from __future__ import annotations

import re
from typing import Iterable, Mapping, Sequence

from docx import Document
from docx.document import Document as DocumentoDocx
from docx.table import Table
from docx.text.paragraph import Paragraph

ETIQUETA = re.compile(r"\{\{\s*([A-Za-z0-9_áéíóúÁÉÍÓÚñÑ]+)\s*\}\}")


def _parrafos_de_tabla(tabla: Table) -> Iterable[Paragraph]:
    for fila in tabla.rows:
        for celda in fila.cells:
            yield from celda.paragraphs
            for anidada in celda.tables:
                yield from _parrafos_de_tabla(anidada)


def recorrer_parrafos(doc: DocumentoDocx) -> Iterable[Paragraph]:
    """Todos los párrafos del documento: cuerpo, tablas, encabezados y pies."""
    yield from doc.paragraphs
    for tabla in doc.tables:
        yield from _parrafos_de_tabla(tabla)
    for seccion in doc.sections:
        for parte in (seccion.header, seccion.footer,
                      seccion.first_page_header, seccion.first_page_footer,
                      seccion.even_page_header, seccion.even_page_footer):
            if parte is None:
                continue
            yield from parte.paragraphs
            for tabla in parte.tables:
                yield from _parrafos_de_tabla(tabla)


def etiquetas_de_parrafo(parrafo: Paragraph) -> list[str]:
    return ETIQUETA.findall(parrafo.text)


def etiquetas_del_documento(doc: DocumentoDocx) -> list[str]:
    """Etiquetas presentes en la plantilla, sin repetir y en orden de aparición."""
    vistas: dict[str, None] = {}
    for parrafo in recorrer_parrafos(doc):
        for clave in etiquetas_de_parrafo(parrafo):
            vistas.setdefault(clave, None)
    return list(vistas)


def _mapa_de_runs(parrafo: Paragraph) -> tuple[str, list[tuple[int, int, int]]]:
    """Texto completo del párrafo y, por run, su rango [inicio, fin) en ese texto."""
    texto = []
    rangos = []
    cursor = 0
    for indice, run in enumerate(parrafo.runs):
        contenido = run.text or ""
        rangos.append((cursor, cursor + len(contenido), indice))
        texto.append(contenido)
        cursor += len(contenido)
    return "".join(texto), rangos


def reemplazar_tramos(parrafo: Paragraph,
                      tramos: Sequence[tuple[int, int, str]]) -> bool:
    """Cambia tramos del texto del párrafo conservando el formato.

    Cada tramo es (inicio, fin, reemplazo) sobre el texto completo del párrafo.
    El reemplazo se escribe en el run donde empieza el tramo y se vacía lo que
    ocupaba en los siguientes, que es lo que mantiene la negrita, la fuente y
    el interlineado del original.

    Es la misma operación que necesita rellenar una etiqueta y que convertir un
    espacio en blanco en etiqueta, así que vive aquí una sola vez.
    """
    if not tramos:
        return False
    texto, rangos = _mapa_de_runs(parrafo)
    nuevo = {indice: texto[ini:fin] for ini, fin, indice in rangos}

    for ini, fin, reemplazo in sorted(tramos, reverse=True):
        for r_ini, r_fin, indice in rangos:
            if r_fin <= ini or r_ini >= fin:
                continue
            local_ini = max(ini, r_ini) - r_ini
            local_fin = min(fin, r_fin) - r_ini
            fragmento = nuevo[indice]
            sustituto = reemplazo if r_ini <= ini < r_fin else ""
            nuevo[indice] = (fragmento[:local_ini] + sustituto
                             + fragmento[local_fin:])

    cambiado = False
    for indice, run in enumerate(parrafo.runs):
        if run.text != nuevo[indice]:
            run.text = nuevo[indice]
            cambiado = True
    return cambiado


def _reemplazar_en_parrafo(parrafo: Paragraph, datos: Mapping[str, str],
                           faltantes: set[str]) -> None:
    texto = parrafo.text
    if "{{" not in texto:
        return

    tramos: list[tuple[int, int, str]] = []
    for coincidencia in ETIQUETA.finditer(texto):
        clave = coincidencia.group(1)
        if clave not in datos:
            faltantes.add(clave)
            continue
        valor = "" if datos[clave] is None else str(datos[clave])
        tramos.append((*coincidencia.span(), valor))
    reemplazar_tramos(parrafo, tramos)


def rellenar(doc: DocumentoDocx, datos: Mapping[str, str]) -> set[str]:
    """Reemplaza las etiquetas del documento. Devuelve las que no tenían dato."""
    faltantes: set[str] = set()
    for parrafo in recorrer_parrafos(doc):
        _reemplazar_en_parrafo(parrafo, datos, faltantes)
    return faltantes


def abrir(ruta_o_archivo) -> DocumentoDocx:
    return Document(ruta_o_archivo)
