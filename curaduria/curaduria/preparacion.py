"""Convierte una plantilla real de la curaduría en plantilla de la app.

Hace de una sola pasada las tres cosas que si no habría que hacer a mano:

1. Cambia los espacios en blanco (`XXXX`, `____`, `XX de ___ de 2026`) por la
   etiqueta que corresponde.
2. Cambia los datos fijos del despacho que están escritos literalmente (el
   nombre de la curadora, el municipio) por su etiqueta, para que salgan de la
   configuración.
3. Recorta el catálogo de párrafos y deja en su lugar las anclas de sección.
   Los párrafos no se pierden: ya están en la biblioteca de bloques, y desde
   ahí se escogen según el caso. Es justamente lo que evita borrar media
   resolución cada vez.

Siempre escribe un archivo nuevo. La plantilla original no se toca.
"""

from __future__ import annotations

import unicodedata
from dataclasses import dataclass, field

from docx.document import Document as DocumentoDocx
from docx.text.paragraph import Paragraph

from .auditoria import Cambio, Hallazgo, convertir, etiquetas_conocidas
from .plantilla import reemplazar_tramos

# Marcas que delimitan las secciones del acto, tal como están escritas.
MARCA_ENCABEZADO = "EN EJERCICIO DE LAS FUNCIONES"
MARCA_CONSIDERANDO = "CONSIDERANDO"
MARCA_RESUELVE = "RESUELVE"
MARCAS_CIERRE = ("NOTIFIQUESE", "NOTIFÍQUESE")


def _plano(texto: str) -> str:
    """Mayúsculas y sin tildes, para comparar marcas sin depender del acento."""
    sin_tildes = "".join(
        c for c in unicodedata.normalize("NFD", texto)
        if unicodedata.category(c) != "Mn")
    return " ".join(sin_tildes.upper().split())


@dataclass
class Informe:
    etiquetas_puestas: list[Cambio] = field(default_factory=list)
    literales: list[Cambio] = field(default_factory=list)
    anclas: list[str] = field(default_factory=list)
    parrafos_recortados: int = 0
    pendientes: list[Hallazgo] = field(default_factory=list)
    nota_estructura: str = ""

    @property
    def total_cambios(self) -> int:
        return len(self.etiquetas_puestas) + len(self.literales)

    def resumen(self) -> str:
        partes = [f"{len(self.etiquetas_puestas)} espacios convertidos en "
                  "etiquetas"]
        if self.literales:
            partes.append(f"{len(self.literales)} datos del despacho")
        if self.anclas:
            partes.append(f"{len(self.anclas)} anclas puestas")
        if self.parrafos_recortados:
            partes.append(f"{self.parrafos_recortados} párrafos del catálogo "
                          "recortados")
        if self.pendientes:
            partes.append(f"{len(self.pendientes)} espacios sin convertir")
        return " · ".join(partes)


# --- Datos fijos escritos literalmente -------------------------------------

def reemplazar_literales(doc: DocumentoDocx,
                         valores: dict[str, str]) -> list[Cambio]:
    """Cambia valores literales por su etiqueta: "Ibagué" -> {{municipio}}.

    Solo se tocan textos que no estén ya dentro de una etiqueta, y se van del
    más largo al más corto para que el nombre completo gane sobre el apellido.
    """
    cambios: list[Cambio] = []
    pares = sorted(((v.strip(), k) for k, v in valores.items()
                    if v and len(v.strip()) >= 4),
                   key=lambda par: -len(par[0]))
    if not pares:
        return cambios

    for indice, parrafo in enumerate(doc.paragraphs, start=1):
        texto = parrafo.text
        if not texto.strip():
            continue
        tramos: list[tuple[int, int, str]] = []
        ocupado: list[tuple[int, int]] = []
        for valor, etiqueta in pares:
            desde = 0
            while True:
                posicion = texto.find(valor, desde)
                if posicion < 0:
                    break
                fin = posicion + len(valor)
                desde = fin
                if any(ini < fin and posicion < f for ini, f in ocupado):
                    continue
                if "{{" in texto[max(0, posicion - 2):fin + 2]:
                    continue
                ocupado.append((posicion, fin))
                tramos.append((posicion, fin, f"{{{{{etiqueta}}}}}"))
                cambios.append(Cambio(parrafo=indice, original=valor,
                                      etiqueta=etiqueta,
                                      contexto=texto[max(0, posicion - 40):fin]))
        reemplazar_tramos(parrafo, tramos)
    return cambios


# --- Esqueleto con anclas ---------------------------------------------------

def _indice_de_marca(parrafos: list[Paragraph], marcas,
                     desde: int = 0) -> int | None:
    """Busca el párrafo del título de una sección.

    Los títulos van solos en su párrafo ("CONSIDERANDO", "RESUELVE"), así que
    se exige que el párrafo sea corto y empiece por la marca. Si se buscara la
    marca en cualquier parte, "se resuelve dicha intervención" —que es un
    considerando— pasaría por el título RESUELVE y cortaría el documento por
    la mitad.
    """
    if isinstance(marcas, str):
        marcas = (marcas,)
    planas = [_plano(m) for m in marcas]
    for indice in range(desde, len(parrafos)):
        plano = _plano(parrafos[indice].text)
        if not plano or len(plano) > 45:
            continue
        if any(plano.startswith(marca) for marca in planas):
            return indice
    return None


def _poner_ancla(parrafos: list[Paragraph], desde: int, hasta: int,
                 ancla: str) -> int:
    """Deja un solo párrafo con {{ancla}} donde había muchos. Devuelve cuántos
    párrafos recortó."""
    if desde > hasta:
        return 0
    referencia = parrafos[desde]
    reemplazar_tramos(referencia, [(0, len(referencia.text),
                                    f"{{{{{ancla}}}}}")])
    if not referencia.text.strip():
        for run in list(referencia.runs):
            run._element.getparent().remove(run._element)
        referencia.add_run(f"{{{{{ancla}}}}}")
    recortados = 0
    for parrafo in parrafos[desde + 1:hasta + 1]:
        if parrafo.text.strip():
            recortados += 1
        parrafo._p.getparent().remove(parrafo._p)
    return recortados


def poner_anclas(doc: DocumentoDocx) -> tuple[list[str], int, str]:
    """Recorta el catálogo y deja las anclas. (anclas, recortados, nota)."""
    parrafos = list(doc.paragraphs)
    indice_considerando = _indice_de_marca(parrafos, MARCA_CONSIDERANDO)
    if indice_considerando is None:
        return [], 0, ("No encontré el título CONSIDERANDO: dejé la plantilla "
                       "sin anclas. Agréguelas a mano donde deban ir.")
    indice_resuelve = _indice_de_marca(parrafos, MARCA_RESUELVE,
                                       indice_considerando + 1)
    if indice_resuelve is None:
        return [], 0, ("No encontré el título RESUELVE después de "
                       "CONSIDERANDO: dejé la plantilla sin anclas.")
    indice_cierre = _indice_de_marca(parrafos, MARCAS_CIERRE,
                                     indice_resuelve + 1)

    anclas: list[str] = []
    recortados = 0

    # De atrás hacia adelante: así los índices de arriba no se mueven.
    if indice_cierre is not None:
        recortados += _poner_ancla(parrafos, indice_cierre,
                                   len(parrafos) - 1, "cierre")
        anclas.append("cierre")
        fin_resuelve = indice_cierre - 1
    else:
        fin_resuelve = len(parrafos) - 1

    recortados += _poner_ancla(parrafos, indice_resuelve + 1, fin_resuelve,
                               "resuelve")
    anclas.append("resuelve")

    recortados += _poner_ancla(parrafos, indice_considerando + 1,
                               indice_resuelve - 1, "considerandos")
    anclas.append("considerandos")

    # El preámbulo puede venir partido en varios párrafos: el ancla cubre
    # desde donde empieza hasta justo antes del título CONSIDERANDO.
    indice_encabezado = None
    for indice in range(indice_considerando):
        if MARCA_ENCABEZADO in _plano(parrafos[indice].text):
            indice_encabezado = indice
            break
    if indice_encabezado is not None:
        recortados += _poner_ancla(parrafos, indice_encabezado,
                                   indice_considerando - 1, "encabezado")
        anclas.append("encabezado")

    anclas.reverse()
    return anclas, recortados, ""


# --- Todo junto -------------------------------------------------------------

def preparar(doc: DocumentoDocx, valores_despacho: dict[str, str] | None = None,
             con_anclas: bool = True) -> Informe:
    """Deja la plantilla lista para la app. Modifica el documento recibido."""
    informe = Informe()
    conocidas = etiquetas_conocidas()

    cambios, pendientes = convertir(doc, aplicar=True, conocidas=conocidas)
    informe.etiquetas_puestas = cambios
    informe.pendientes = pendientes

    if valores_despacho:
        informe.literales = reemplazar_literales(doc, valores_despacho)

    if con_anclas:
        anclas, recortados, nota = poner_anclas(doc)
        informe.anclas = anclas
        informe.parrafos_recortados = recortados
        informe.nota_estructura = nota

    return informe
