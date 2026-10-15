"""Genera una plantilla maestra de ejemplo con todas las etiquetas disponibles.

    python crear_plantilla_ejemplo.py [ruta_salida]

Sirve como punto de partida: ábrala en Word, pegue el texto real de sus
resoluciones y conserve las etiquetas {{...}} donde van los datos variables.
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

DESTINO_POR_DEFECTO = Path(__file__).parent / "plantillas" / "plantilla_maestra.docx"


def _p(doc, texto, *, negrita=False, centrado=False, tamano=11):
    parrafo = doc.add_paragraph()
    if centrado:
        parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        parrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = parrafo.add_run(texto)
    run.bold = negrita
    run.font.size = Pt(tamano)
    return parrafo


def construir() -> Document:
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Arial"
    estilo.font.size = Pt(11)

    encabezado = doc.sections[0].header.paragraphs[0]
    encabezado.alignment = WD_ALIGN_PARAGRAPH.CENTER
    encabezado.add_run("{{curaduria}}").bold = True

    _p(doc, "RESOLUCIÓN No. {{resolucion_numero}}", negrita=True, centrado=True)
    _p(doc, "({{fecha_larga}})", centrado=True)
    _p(doc, "")
    _p(doc, "“Por medio de la cual se resuelve la solicitud de {{tipo_tramite}} "
            "radicada bajo el número {{radicado}} del {{fecha_radicado_larga}}”",
       negrita=True, centrado=True)
    _p(doc, "")
    _p(doc, "{{Articulo}} {{curador_tratamiento}} de {{municipio}}, en ejercicio de "
            "las facultades legales, y")
    _p(doc, "")
    _p(doc, "CONSIDERANDO", negrita=True, centrado=True)
    _p(doc, "")
    _p(doc, "PRIMERO: Que {{identificacion}}, {{verbo_obrar}} {{calidad}} "
            "{{documento}} y {{verbo_solicitar}} licencia para el predio ubicado en "
            "{{direccion}}, barrio {{barrio}}, del municipio de {{municipio}} "
            "({{departamento}}).")
    _p(doc, "")
    _p(doc, "SEGUNDO: Que el predio se identifica con matrícula inmobiliaria No. "
            "{{matricula}} y cédula catastral No. {{cedula_catastral}}, con un área "
            "de lote de {{area_lote_letras}} y un área construida de "
            "{{area_construida_letras}}.")
    _p(doc, "")
    _p(doc, "TERCERO: Que el predio se encuentra en tratamiento urbanístico "
            "{{tratamiento_urbanistico}}, con uso {{uso}} y una altura de "
            "{{pisos}} ({{pisos_cifra}}) pisos.")
    _p(doc, "")
    _p(doc, "CUARTO: Que los planos fueron elaborados por {{proyectista}}, con "
            "matrícula profesional No. {{matricula_profesional}}.")
    _p(doc, "")
    _p(doc, "RESUELVE", negrita=True, centrado=True)
    _p(doc, "")
    _p(doc, "ARTÍCULO PRIMERO: {{sentido}} {{tipo_tramite}} solicitada por "
            "{{articulo}} {{propietario}} {{nombres}} ({{documentos}}), para el "
            "predio ubicado en {{direccion}} de {{municipio}}.")
    _p(doc, "")
    _p(doc, "ARTÍCULO SEGUNDO: La presente licencia tendrá una vigencia de "
            "{{vigencia}}, contados a partir de la fecha de su ejecutoria.")
    _p(doc, "")
    _p(doc, "ARTÍCULO TERCERO: Notificar personalmente el contenido de esta "
            "resolución {{contraccion_a}} {{interesado}} {{nombres}}, {{quien}} "
            "{{verbo_podra}} interponer recurso de reposición y apelación dentro de "
            "los diez (10) días siguientes a la notificación.")
    _p(doc, "")
    _p(doc, "ARTÍCULO CUARTO: {{observaciones}}")
    _p(doc, "")
    _p(doc, "Dada en {{municipio}}, a los {{fecha_letras}}.", centrado=True)
    _p(doc, "")
    _p(doc, "")
    _p(doc, "NOTIFÍQUESE Y CÚMPLASE", negrita=True, centrado=True)
    _p(doc, "")
    _p(doc, "")
    _p(doc, "{{curador}}", negrita=True, centrado=True)
    _p(doc, "{{curador_tratamiento}} de {{municipio}}", centrado=True)

    pie = doc.sections[0].footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pie.add_run("Resolución {{resolucion_numero}} del {{fecha}} — Radicado {{radicado}}")

    return doc


def main() -> None:
    destino = Path(sys.argv[1]) if len(sys.argv) > 1 else DESTINO_POR_DEFECTO
    destino.parent.mkdir(parents=True, exist_ok=True)
    construir().save(destino)
    print(f"Plantilla de ejemplo creada en: {destino}")


if __name__ == "__main__":
    main()
